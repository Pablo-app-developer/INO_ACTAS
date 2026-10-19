"""Construye el acta en Word (.docx) a partir de `motor/plantillas/acta_ino.docx`.

De la plantilla se conservan la página, los márgenes, los estilos y la numeración (viñetas);
el contenido de ejemplo se descarta y se reemplaza. El formato de párrafos y tablas reproduce
el del acta 030 del INO: Arial Narrow, texto justificado, viñetas con etiqueta en negrita y
tablas sin bordes. Si cambia el formato del acta, se ajusta en las constantes de abajo.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from .acta import Acta, Lista, Parrafo, Subtitulo, Tabla

PLANTILLA_ACTA = Path(__file__).parent / "plantillas" / "acta_ino.docx"

FUENTE = "Arial Narrow"
TAMANO_TITULO = Pt(16)
NUM_VINETA = 1                     # numId de viñetas (•) en la numeración de la plantilla
ANCHO_UTIL = 8838                  # twips: página carta menos márgenes de la plantilla
ANCHOS_TABLA = {2: [2217, 6621], 3: [1918, 5916, 1004]}   # igual que las tablas del acta
TITULO_COMPROMISOS = "Tabla de Compromisos y Tareas Pendientes"
TITULO_GENERAL = "Información General y Contexto Institucional"


def _parrafo(doc, texto="", *, negrita=False, centrado=False, justificado=True, tamano=None,
             vineta=False) -> Paragraph:
    p = doc.add_paragraph()
    if centrado:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif justificado:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if vineta:
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = NUM_VINETA
    if texto:
        _run(p, texto, negrita=negrita, tamano=tamano)
    return p


def _run(p: Paragraph, texto: str, *, negrita=False, tamano=None):
    r = p.add_run(texto)
    r.font.name = FUENTE
    if negrita:
        r.bold = True
    if tamano:
        r.font.size = tamano
    return r


def _punto(doc, etiqueta: str, texto: str) -> None:
    p = _parrafo(doc, vineta=True)
    etiqueta = etiqueta.strip()
    if etiqueta:
        _run(p, etiqueta, negrita=True)
        _run(p, f" {texto.strip()}")
    else:
        _run(p, texto.strip())


def _tabla(doc, columnas: list[str], filas: list[list[str]]) -> None:
    n = len(columnas)
    anchos = ANCHOS_TABLA.get(n) or [ANCHO_UTIL // n] * n
    xml = (
        f"<w:tbl {nsdecls('w')}><w:tblPr><w:tblW w:w=\"0\" w:type=\"auto\"/>"
        "<w:tblCellMar><w:top w:w=\"15\" w:type=\"dxa\"/><w:left w:w=\"15\" w:type=\"dxa\"/>"
        "<w:bottom w:w=\"15\" w:type=\"dxa\"/><w:right w:w=\"15\" w:type=\"dxa\"/></w:tblCellMar>"
        "<w:tblLook w:val=\"04A0\" w:firstRow=\"1\" w:lastRow=\"0\" w:firstColumn=\"1\" "
        "w:lastColumn=\"0\" w:noHBand=\"0\" w:noVBand=\"1\"/></w:tblPr><w:tblGrid>"
        + "".join(f"<w:gridCol w:w=\"{a}\"/>" for a in anchos)
        + "</w:tblGrid></w:tbl>"
    )
    tbl = parse_xml(xml)
    doc.element.body.sectPr.addprevious(tbl)

    def fila(celdas: list[str], encabezado: bool) -> None:
        tr = parse_xml(f"<w:tr {nsdecls('w')}/>")
        for i, texto in enumerate(celdas):
            tc = parse_xml(f"<w:tc {nsdecls('w')}><w:tcPr><w:tcW w:w=\"0\" w:type=\"auto\"/></w:tcPr></w:tc>")
            p_xml = parse_xml(f"<w:p {nsdecls('w')}/>")
            tc.append(p_xml)
            parrafo = Paragraph(p_xml, None)
            # Encabezado y primera columna en negrita, como en el acta
            _run(parrafo, texto, negrita=encabezado or i == 0)
            tr.append(tc)
        tbl.append(tr)

    fila(columnas, True)
    for f in filas:
        fila((f + [""] * n)[:n], False)   # tolera filas cortas o largas del modelo


def _vaciar_cuerpo(doc) -> None:
    cuerpo = doc.element.body
    for hijo in list(cuerpo):
        if hijo.tag != qn("w:sectPr"):
            cuerpo.remove(hijo)


def construir_docx(acta: Acta, destino: Path, plantilla: Path = PLANTILLA_ACTA) -> Path:
    doc = Document(str(plantilla))
    _vaciar_cuerpo(doc)

    _parrafo(doc, f"ACTA {acta.numero}".strip(), negrita=True, centrado=True, tamano=TAMANO_TITULO)
    _parrafo(doc)
    _parrafo(doc, f"Acta de Reunión: {acta.titulo_reunion} - {acta.fecha}", negrita=True, justificado=False)

    # 1. Información general (sección fija)
    _parrafo(doc, f"1. {TITULO_GENERAL}", negrita=True)
    if acta.introduccion.strip():
        _parrafo(doc, acta.introduccion.strip())
    _parrafo(doc, "Detalles de la sesión:", negrita=True)
    _punto(doc, "Fecha:", f"{acta.fecha}.")
    if acta.hora.strip():
        _punto(doc, "Hora:", acta.hora.strip())
    _punto(doc, "Objetivo de la sesión:", acta.objetivo)
    if acta.participantes:
        _punto(doc, "Participantes:", ", ".join(acta.participantes) + ".")
    if acta.apertura.strip():
        _parrafo(doc, acta.apertura.strip())

    # 2..N. Secciones temáticas
    for n, seccion in enumerate(acta.secciones, start=2):
        _parrafo(doc, f"{n}. {seccion.titulo.strip()}", negrita=True)
        for bloque in seccion.bloques:
            if isinstance(bloque, Parrafo):
                _parrafo(doc, bloque.texto.strip())
            elif isinstance(bloque, Subtitulo):
                _parrafo(doc, bloque.texto.strip(), negrita=True)
            elif isinstance(bloque, Lista):
                for punto in bloque.puntos:
                    _punto(doc, punto.etiqueta, punto.texto)
            elif isinstance(bloque, Tabla) and bloque.columnas:
                _tabla(doc, bloque.columnas, bloque.filas)
                _parrafo(doc)

    # Última sección: compromisos
    _parrafo(doc, f"{len(acta.secciones) + 2}. {TITULO_COMPROMISOS}", negrita=True, justificado=False)
    if acta.compromisos:
        _tabla(doc, ["Responsable", "Compromiso / Tarea", "Prioridad"],
               [[c.responsable, c.compromiso, c.prioridad] for c in acta.compromisos])
    else:
        _parrafo(doc, "No se registraron compromisos en esta sesión.")
    _parrafo(doc)

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(destino))
    return destino
