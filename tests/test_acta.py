import json
from datetime import date
from types import SimpleNamespace
from unittest import mock

import pytest
from docx import Document
from docx.oxml.ns import qn

from motor.acta import (
    Acta, Compromiso, DatosActa, ErrorActa, Lista, Parrafo, Punto, Seccion, Subtitulo, Tabla,
    fecha_larga, generar_acta,
)
from motor.docx_acta import PLANTILLA_ACTA, construir_docx

JSON_VALIDO = {
    "titulo_reunion": "Observatorio de Datos e Innovación Clínica",
    "objetivo": "Revisar el avance de los proyectos.",
    "participantes": ["Ing. Rafael Ortiz", "Pablo Ramírez"],
    "introduccion": "La sesión se centró en el avance.",
    "apertura": "",
    "secciones": [{"titulo": "Despliegue", "bloques": [{"tipo": "parrafo", "texto": "Se evaluó el despliegue."}]}],
    "compromisos": [{"responsable": "Pablo", "compromiso": "Validar el monitor.", "prioridad": "Alta"}],
    # El modelo no debe poner estos campos; si los pone, se ignoran
    "numero": "999", "fecha": "1 de enero de 2000", "hora": "3:00",
}
DATOS = DatosActa(numero="030", fecha=date(2026, 9, 2), hora="8:30 a. m.")


def respuesta(contenido: str):
    return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=contenido))])


def cliente_falso(*contenidos: str):
    cliente = mock.MagicMock()
    cliente.chat.completions.create.side_effect = [respuesta(c) for c in contenidos]
    return cliente


@pytest.fixture(autouse=True)
def con_clave(monkeypatch):
    monkeypatch.setenv("DEEPSEEK_API_KEY", "falsa")


# ------------------------------------------------------------------ generación

def test_fecha_larga():
    assert fecha_larga(date(2026, 9, 2)) == "02 de septiembre de 2026"
    assert fecha_larga(date(2026, 12, 31)) == "31 de diciembre de 2026"


def test_generar_acta_usa_los_datos_del_usuario_y_no_los_del_modelo():
    cliente = cliente_falso(json.dumps(JSON_VALIDO))
    with mock.patch("motor.acta.OpenAI", return_value=cliente):
        acta = generar_acta("texto", "reunion", DATOS, "contexto X")
    assert (acta.numero, acta.fecha, acta.hora) == ("030", "02 de septiembre de 2026", "8:30 a. m.")
    assert acta.compromisos[0].prioridad == "Alta"
    mensajes = cliente.chat.completions.create.call_args.kwargs["messages"]
    assert "contexto X" in mensajes[1]["content"] and "texto" in mensajes[1]["content"]
    assert cliente.chat.completions.create.call_args.kwargs["response_format"] == {"type": "json_object"}


def test_reintenta_una_vez_si_el_json_es_invalido():
    cliente = cliente_falso("esto no es json", json.dumps(JSON_VALIDO))
    with mock.patch("motor.acta.OpenAI", return_value=cliente):
        acta = generar_acta("t", "r", DATOS)
    assert cliente.chat.completions.create.call_count == 2
    assert acta.titulo_reunion.startswith("Observatorio")


def test_falla_si_el_modelo_no_cumple_el_esquema_dos_veces():
    sin_secciones = json.dumps({"titulo_reunion": "x"})
    cliente = cliente_falso(sin_secciones, sin_secciones)
    with mock.patch("motor.acta.OpenAI", return_value=cliente), pytest.raises(ErrorActa):
        generar_acta("t", "r", DATOS)


def test_prioridad_invalida_se_rechaza():
    malo = {**JSON_VALIDO, "compromisos": [{"responsable": "a", "compromiso": "b", "prioridad": "Urgentísima"}]}
    cliente = cliente_falso(json.dumps(malo), json.dumps(malo))
    with mock.patch("motor.acta.OpenAI", return_value=cliente), pytest.raises(ErrorActa):
        generar_acta("t", "r", DATOS)


def test_sin_clave_no_se_puede(monkeypatch):
    monkeypatch.delenv("DEEPSEEK_API_KEY")
    with pytest.raises(ErrorActa, match="DEEPSEEK_API_KEY"):
        generar_acta("t", "r", DATOS)


# ------------------------------------------------------------------------ docx

def acta_ejemplo(**cambios) -> Acta:
    base = dict(
        numero="031", titulo_reunion="Observatorio de Datos", fecha="09 de septiembre de 2026",
        hora="9:00 a. m.", objetivo="Revisar avances.", participantes=["Ana", "Luis"],
        introduccion="Introducción de la sesión.", apertura="Se abrió la sesión.",
        secciones=[
            Seccion(titulo="Arquitectura", bloques=[
                Parrafo(texto="Texto de la sección."),
                Subtitulo(texto="Resumen"),
                Lista(puntos=[Punto(etiqueta="Estado:", texto="Cargado."), Punto(texto="Sin etiqueta.")]),
                Tabla(columnas=["Componente", "Decisión"], filas=[["VPS", "Externo"], ["Nginx", "Proxy"]]),
            ]),
            Seccion(titulo="Datos de sueño", bloques=[Parrafo(texto="Otro tema.")]),
        ],
        compromisos=[Compromiso(responsable="Ana", compromiso="Enviar el informe.", prioridad="Media"),
                     Compromiso(responsable="Luis", compromiso="Revisar el monitor.", prioridad="Recurrente")],
    )
    base.update(cambios)
    return Acta(**base)


def textos(doc):
    return [p.text for p in doc.paragraphs]


def test_docx_estructura_y_orden(tmp_path):
    destino = construir_docx(acta_ejemplo(), tmp_path / "acta.docx")
    doc = Document(str(destino))
    t = textos(doc)

    assert t[0] == "ACTA 031"
    assert t[2] == "Acta de Reunión: Observatorio de Datos - 09 de septiembre de 2026"
    orden = [x for x in t if x[:2] in ("1.", "2.", "3.", "4.")]
    assert orden == ["1. Información General y Contexto Institucional", "2. Arquitectura",
                     "3. Datos de sueño", "4. Tabla de Compromisos y Tareas Pendientes"]
    assert "Fecha: 09 de septiembre de 2026." in t
    assert "Hora: 9:00 a. m." in t
    assert "Participantes: Ana, Luis." in t
    assert "Estado: Cargado." in t and "Sin etiqueta." in t
    # Nada del contenido de ejemplo de la plantilla
    assert not any("ERS" in x or "Ergoespirometría" in x for x in t)


def test_docx_formato_igual_al_del_acta(tmp_path):
    doc = Document(str(construir_docx(acta_ejemplo(), tmp_path / "acta.docx")))
    titulo = doc.paragraphs[0]
    assert titulo.alignment == 1                                    # centrado
    assert titulo.runs[0].bold and titulo.runs[0].font.size.pt == 16
    cuerpo = next(p for p in doc.paragraphs if p.text == "Introducción de la sesión.")
    assert cuerpo.alignment == 3                                    # justificado
    assert all(r.font.name == "Arial Narrow" for p in doc.paragraphs for r in p.runs)

    fecha = next(p for p in doc.paragraphs if p.text.startswith("Fecha:"))
    assert fecha.runs[0].text == "Fecha:" and fecha.runs[0].bold    # etiqueta en negrita
    assert not fecha.runs[1].bold
    num_pr = fecha._p.pPr.numPr
    assert num_pr is not None and num_pr.numId.val == 1             # viñeta de la plantilla


def test_docx_tablas(tmp_path):
    doc = Document(str(construir_docx(acta_ejemplo(), tmp_path / "acta.docx")))
    assert len(doc.tables) == 2
    tecnica, compromisos = doc.tables
    assert [c.text for c in tecnica.rows[0].cells] == ["Componente", "Decisión"]
    assert [c.text for c in tecnica.rows[2].cells] == ["Nginx", "Proxy"]
    assert [c.text for c in compromisos.rows[0].cells] == ["Responsable", "Compromiso / Tarea", "Prioridad"]
    assert [c.text for c in compromisos.rows[2].cells] == ["Luis", "Revisar el monitor.", "Recurrente"]
    assert compromisos.rows[0].cells[0].paragraphs[0].runs[0].bold
    grid = [int(g.get(qn("w:w"))) for g in compromisos._tbl.tblGrid.findall(qn("w:gridCol"))]
    assert grid == [1918, 5916, 1004]                               # mismos anchos que el acta 030


def test_docx_hora_opcional_y_sin_compromisos(tmp_path):
    doc = Document(str(construir_docx(acta_ejemplo(hora="", apertura="", compromisos=[]), tmp_path / "a.docx")))
    t = textos(doc)
    assert not any(x.startswith("Hora:") for x in t)
    assert "No se registraron compromisos en esta sesión." in t
    assert len(doc.tables) == 1                                     # solo la de la sección


def test_docx_filas_irregulares_no_rompen(tmp_path):
    acta = acta_ejemplo(secciones=[Seccion(titulo="T", bloques=[
        Tabla(columnas=["A", "B"], filas=[["solo uno"], ["1", "2", "sobra"]])])])
    doc = Document(str(construir_docx(acta, tmp_path / "a.docx")))
    assert [c.text for c in doc.tables[0].rows[1].cells] == ["solo uno", ""]
    assert [c.text for c in doc.tables[0].rows[2].cells] == ["1", "2"]


def test_docx_conserva_pagina_de_la_plantilla(tmp_path):
    original = Document(str(PLANTILLA_ACTA)).sections[0]
    nuevo = Document(str(construir_docx(acta_ejemplo(), tmp_path / "a.docx"))).sections[0]
    assert (nuevo.page_width, nuevo.left_margin, nuevo.top_margin) == (
        original.page_width, original.left_margin, original.top_margin)
